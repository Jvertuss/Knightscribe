from io import BytesIO
from pathlib import Path

from pypdf import PdfReader
from pptx import Presentation


MAX_CHUNK_CHARS = 1200
CHUNK_OVERLAP_CHARS = 200


def split_text_into_chunks(text: str) -> list[str]:
    """
    Split extracted text into smaller overlapping chunks
    suitable for later AI retrieval.
    """

    cleaned_text = " ".join(text.split())

    if not cleaned_text:
        return []

    chunks = []
    start = 0

    while start < len(cleaned_text):
        end = min(
            start + MAX_CHUNK_CHARS,
            len(cleaned_text),
        )

        # Try not to split in the middle of a word.
        if end < len(cleaned_text):
            last_space = cleaned_text.rfind(
                " ",
                start,
                end,
            )

            if last_space > start + (MAX_CHUNK_CHARS // 2):
                end = last_space

        chunk = cleaned_text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end >= len(cleaned_text):
            break

        start = max(
            end - CHUNK_OVERLAP_CHARS,
            start + 1,
        )

    return chunks


def extract_pdf(data: bytes) -> list[dict]:
    """
    Extract text from a PDF while preserving page numbers.
    """

    reader = PdfReader(BytesIO(data))

    extracted_chunks = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        page_text = page.extract_text() or ""

        chunks = split_text_into_chunks(page_text)

        for chunk in chunks:
            extracted_chunks.append(
                {
                    "content": chunk,
                    "source_location": f"Page {page_number}",
                }
            )

    return extracted_chunks


def extract_powerpoint(data: bytes) -> list[dict]:
    """
    Extract visible text from PowerPoint slides while
    preserving slide numbers.
    """

    presentation = Presentation(BytesIO(data))

    extracted_chunks = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_text_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    slide_text_parts.append(text)

        slide_text = "\n".join(slide_text_parts)

        chunks = split_text_into_chunks(slide_text)

        for chunk in chunks:
            extracted_chunks.append(
                {
                    "content": chunk,
                    "source_location": f"Slide {slide_number}",
                }
            )

    return extracted_chunks


def extract_text_file(data: bytes) -> list[dict]:
    """
    Extract text from a plain-text file.
    """

    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        raise ValueError(
            "Text file must use UTF-8 encoding"
        )

    chunks = split_text_into_chunks(text)

    return [
        {
            "content": chunk,
            "source_location": "Text file",
        }
        for chunk in chunks
    ]


def extract_document(
    filename: str,
    data: bytes,
) -> tuple[str, list[dict]]:
    """
    Detect a supported document type and extract
    searchable chunks from it.

    Returns:
        document_type
        chunks
    """

    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        document_type = "pdf"
        chunks = extract_pdf(data)

    elif extension == ".pptx":
        document_type = "powerpoint"
        chunks = extract_powerpoint(data)

    elif extension == ".txt":
        document_type = "text"
        chunks = extract_text_file(data)

    else:
        raise ValueError(
            "Unsupported file type. "
            "KnightScribe currently supports PDF, PPTX, and TXT files."
        )

    if not chunks:
        raise ValueError(
            "No readable text could be extracted from this file."
        )

    return document_type, chunks